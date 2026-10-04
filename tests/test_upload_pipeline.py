import io
import unittest
import os
import sys

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from app import create_app
import pymupdf
import docx

class UploadPipelineTestCase(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = create_app()
        cls.app.config["TESTING"] = True
        cls.client = cls.app.test_client()
        with cls.client.session_transaction() as sess:
            sess["user_id"] = 1
            sess["user_name"] = "alex rivera"
            sess["user_email"] = "demo@jobalign.ai"

    def test_upload_txt_resume(self):
        txt_content = (
            "marcus vance\n"
            "marcus.vance@techreach.io | (555) 777-8899 | austin, tx\n"
            "summary\n"
            "lead software engineer with 5 years experience in python, django, mysql, and kubernetes.\n"
            "technical skills: python, django, postgresql, mysql, docker, kubernetes, aws, redis\n"
            "education: b.s. in computer engineering, texas a&m university, 2019\n"
            "experience: lead backend engineer at cloudbridge tech (2019 - present)\n"
            "- managed distributed postgresql and mysql data systems\n"
            "- containerized applications using docker and kubernetes\n"
            "projects: cloud orchestration framework built with python and docker\n"
            "certifications: certified kubernetes administrator"
        )
        data = {
            "resume_file": (io.BytesIO(txt_content.encode("utf-8")), "marcus_resume.txt")
        }
        res = self.client.post("/api/resume/upload", data=data, content_type="multipart/form-data")
        self.assertEqual(res.status_code, 200)
        res_json = res.get_json()
        self.assertTrue(res_json["success"])
        self.assertIn("resume_id", res_json)
        self.assertIn("python", res_json["parsed_data"]["skills"])

    def test_upload_docx_resume(self):
        doc = docx.Document()
        doc.add_heading("elena rostova", 0)
        doc.add_paragraph("elena.rostova@aiinnovate.com | (555) 432-1098 | seattle, wa")
        doc.add_heading("skills", level=1)
        doc.add_paragraph("python, machine learning, deep learning, pytorch, transformers, nlp, scikit-learn, mysql")
        doc.add_heading("education", level=1)
        doc.add_paragraph("master of science in artificial intelligence, university of washington, 2021")
        doc.add_heading("experience", level=1)
        doc.add_paragraph("machine learning engineer at deepsight labs (2021 - present)")
        doc.add_paragraph("trained transformer models and built semantic embedding pipelines.")

        bio = io.BytesIO()
        doc.save(bio)
        bio.seek(0)

        data = {
            "resume_file": (bio, "elena_resume.docx")
        }
        res = self.client.post("/api/resume/upload", data=data, content_type="multipart/form-data")
        self.assertEqual(res.status_code, 200)
        res_json = res.get_json()
        self.assertTrue(res_json["success"])
        self.assertIn("machine learning", res_json["parsed_data"]["skills"])

    def test_upload_pdf_resume(self):
        doc = pymupdf.open()
        page = doc.new_page()
        text = (
            "david chen\n"
            "david.chen@cloudworks.com | (555) 321-6549 | san jose, ca\n"
            "skills: python, flask, mysql, docker, aws, git, rest api\n"
            "education: b.s. in computer science, san jose state university, 2020\n"
            "experience: software engineer at apex data (2020 - present)\n"
            "- built rest apis in python and flask"
        )
        page.insert_text((50, 72), text, fontsize=11)
        pdf_bytes = doc.write()
        doc.close()

        data = {
            "resume_file": (io.BytesIO(pdf_bytes), "david_resume.pdf")
        }
        res = self.client.post("/api/resume/upload", data=data, content_type="multipart/form-data")
        self.assertEqual(res.status_code, 200)
        res_json = res.get_json()
        self.assertTrue(res_json["success"])
        self.assertIn("flask", res_json["parsed_data"]["skills"])

if __name__ == "__main__":
    unittest.main()
