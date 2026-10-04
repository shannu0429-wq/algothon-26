import os
import re
import pymupdf as fitz
import docx
from services.ocr_service import extract_text_from_image, is_ocr_available

def normalize_extracted_text(text):
    if not text:
        return ""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("\u00a0", " ").replace("\u200b", "")
    text = re.sub(r"[•·●▪■►▸◦❖]", " - ", text)
    text = re.sub(r"-\n\s*", "", text)
    text = re.sub(r"[ \t]+", " ", text)
    lines = [line.strip() for line in text.split("\n")]
    cleaned_lines = []
    prev_empty = False
    for line in lines:
        if not line:
            if not prev_empty:
                cleaned_lines.append("")
                prev_empty = True
        else:
            cleaned_lines.append(line)
            prev_empty = False
    return "\n".join(cleaned_lines).strip()

def extract_from_pdf(filepath):
    try:
        doc = fitz.open(filepath)
        full_text = []
        for page_num in range(len(doc)):
            page = doc[page_num]
            page_text = page.get_text("text")
            if page_text and page_text.strip():
                full_text.append(page_text.strip())
        
        extracted = "\n\n".join(full_text).strip()
        if len(extracted) >= 50:
            doc.close()
            return normalize_extracted_text(extracted), None
        
        if is_ocr_available():
            ocr_text = []
            for page_num in range(min(len(doc), 5)):
                page = doc[page_num]
                pix = page.get_pixmap(dpi=200)
                temp_img_path = f"{filepath}_page_{page_num}.png"
                pix.save(temp_img_path)
                p_text, err = extract_text_from_image(temp_img_path)
                if os.path.exists(temp_img_path):
                    try:
                        os.remove(temp_img_path)
                    except Exception:
                        pass
                if p_text:
                    ocr_text.append(p_text)
            doc.close()
            joined_ocr = "\n\n".join(ocr_text).strip()
            if joined_ocr:
                return normalize_extracted_text(joined_ocr), None
            return "", "pdf appears to be a scanned document and no text could be extracted via ocr."
        doc.close()
        if extracted:
            return normalize_extracted_text(extracted), None
        return "", "pdf contains scanned images or no readable text, and ocr engine is not available."
    except Exception as exc:
        return "", f"failed to parse pdf document: {str(exc).lower()}"

def extract_from_docx(filepath):
    try:
        doc = docx.Document(filepath)
        text_parts = []
        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text.strip())
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    text_parts.append(" | ".join(row_cells))
        joined = "\n".join(text_parts)
        if not joined.strip():
            return "", "docx document is empty."
        return normalize_extracted_text(joined), None
    except Exception as exc:
        return "", f"failed to parse docx document: {str(exc).lower()}"

def extract_from_txt(filepath):
    encodings = ["utf-8", "utf-8-sig", "latin-1", "cp1252", "iso-8859-1"]
    for enc in encodings:
        try:
            with open(filepath, "r", encoding=enc) as f:
                content = f.read()
                if content.strip():
                    return normalize_extracted_text(content), None
        except UnicodeDecodeError:
            continue
        except Exception as exc:
            return "", f"failed to read text file: {str(exc).lower()}"
    return "", "unable to decode text file with standard encodings."

def extract_document_text(filepath):
    if not os.path.exists(filepath):
        return "", "file not found on server."
    
    ext = filepath.rsplit(".", 1)[-1].lower() if "." in filepath else ""
    
    if ext == "pdf":
        return extract_from_pdf(filepath)
    elif ext == "docx":
        return extract_from_docx(filepath)
    elif ext == "txt":
        return extract_from_txt(filepath)
    elif ext in ["png", "jpg", "jpeg"]:
        if not is_ocr_available():
            return "", "image resumes require tesseract ocr which is currently unavailable on the system."
        return extract_text_from_image(filepath)
    else:
        return "", f"unsupported file format: .{ext}"
